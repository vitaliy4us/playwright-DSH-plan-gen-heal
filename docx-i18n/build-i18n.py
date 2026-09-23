#!/usr/bin/env python3
"""Build the localised versions of the combined .docx document.

Reuses the markdown converter from C:\\Users\\vital\\AI\\build-docx.py and reads the translated
markdown from docx-i18n/<lang>/. Run from the project root:

    python docx-i18n/build-i18n.py en
    python docx-i18n/build-i18n.py uk
"""
from __future__ import annotations

import importlib.util
import os
import sys
from datetime import date

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

WORKSPACE = r"C:\Users\vital\AI\playwright-DSH-plan-gen-heal"
I18N_ROOT = os.path.join(WORKSPACE, "docx-i18n")
BUILDER_PATH = r"C:\Users\vital\AI\build-docx.py"

# Reuse the converter (runs no code at import: its script body is under a __main__ guard).
_spec = importlib.util.spec_from_file_location("docx_builder", BUILDER_PATH)
builder = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(builder)

# ---------------------------------------------------------------- strings per language
LABELS = {
    "en": {
        "title": "Playwright Test Agents",
        "subtitle": "Course, projects and reports",
        "intro": (
            "A single document compiled from every project material:\n"
            "the Playwright Test Agents course (6 lessons),\n"
            "the GreenKart project (plans, test cases, autotests, defects, agent definitions)\n"
            "and the Conduit project (Playwright MCP vs playwright-cli comparison, generation prompts)."
        ),
        "compiled": "Compiled on {date}",
        "contents": "Contents",
        "parts": [
            (
                "Part I. Course “Playwright Test Agents”",
                "Six lessons: from an introduction to Test Agents to hands-on practice and a comparison "
                "of the manual and agent-driven approaches. The course is adapted to DeepSeek Harness.",
                [
                    ("0. About the course", "course/README.md"),
                    ("1. Introduction to Playwright Test Agents", "course/01-intro-test-agents.md"),
                    ("2. Playwright agents in DeepSeek Harness", "course/02-agents-in-dsh.md"),
                    ("3. How the agents work inside (under the hood)", "course/03-under-the-hood.md"),
                    ("4. Playwright CLI: automation from the command line", "course/04-agent-cli.md"),
                    ("5. Playwright MCP as the base layer", "course/05-playwright-mcp.md"),
                    ("6. Practice: the full cycle and manual vs agents", "course/06-practice-manual-vs-agents.md"),
                ],
            ),
            (
                "Part II. GreenKart project (plan → generate → heal)",
                "A training project on the GreenKart application: test plans, manual test cases, "
                "automated tests, the defects found, and the agent roles for DeepSeek Harness.",
                [
                    ("1. Project overview", "greenkart/README.md"),
                    ("2. Test plan (base)", "greenkart/specs/greenkart.plan.md"),
                    ("3. Cart Edge Cases plan", "greenkart/specs/cart-edge-cases.plan.md"),
                    ("4. Search and Cart plan", "greenkart/specs/search-and-cart.plan.md"),
                    ("5. Manual test plan", "greenkart/test-cases/greenkart-test-plan.md"),
                    ("6. Manual test cases", None),
                    ("7. Defect report BUG-1: the cart is lost when the page is reloaded immediately",
                     "greenkart/BUG-1-cart-is-lost-on-immediate-reload.md"),
                    ("8. Defect report BUG-2: an order can be placed with an empty cart",
                     "greenkart/BUG-2-empty-cart-can-be-ordered.md"),
                    ("9. Agent roles (skills for DSH)", None),
                    ("10. Playwright agent definitions (init-agents)", None),
                    ("11. Project log (SESSION-NOTES)", "greenkart/SESSION-NOTES.md"),
                ],
            ),
            (
                "Part III. Conduit project (Playwright MCP vs playwright-cli)",
                "The same test case is automated twice — through Playwright MCP and through playwright-cli — "
                "in order to compare the two approaches.",
                [
                    ("1. Overview and comparison of the approaches", "conduit/README.md"),
                    ("2. Test plan", "conduit/conduit.plan.md"),
                    ("3. Generation prompts (MCP and CLI)", "conduit/prompts.txt"),
                    ("4. Project log (SESSION-NOTES)", "conduit/SESSION-NOTES.md"),
                ],
            ),
        ],
        "groups": {
            "6. Manual test cases": ("greenkart/test-cases", "TC-"),
            "9. Agent roles (skills for DSH)": ("greenkart/skills", None),
            "10. Playwright agent definitions (init-agents)": ("greenkart/agents", None),
        },
    },
    "uk": {
        "title": "Playwright Test Agents",
        "subtitle": "Курс, проєкти та звіти",
        "intro": (
            "Єдиний документ, укладений з усіх матеріалів проєкту:\n"
            "навчальний курс із Playwright Test Agents (6 уроків),\n"
            "проєкт GreenKart (плани, тест-кейси, автотести, дефекти, визначення агентів)\n"
            "і проєкт Conduit (порівняння Playwright MCP і playwright-cli, промпти генерації)."
        ),
        "compiled": "Укладено {date}",
        "contents": "Зміст",
        "parts": [
            (
                "Частина I. Навчальний курс «Playwright Test Agents»",
                "Шість уроків: від вступу до Test Agents до практики та порівняння ручного й агентного "
                "підходів. Курс адаптовано під роботу в DeepSeek Harness.",
                [
                    ("0. Про курс", "course/README.md"),
                    ("1. Вступ до Playwright Test Agents", "course/01-intro-test-agents.md"),
                    ("2. Агенти Playwright у DeepSeek Harness", "course/02-agents-in-dsh.md"),
                    ("3. Як агенти влаштовані всередині (under the hood)", "course/03-under-the-hood.md"),
                    ("4. Playwright CLI: автоматизація з командного рядка", "course/04-agent-cli.md"),
                    ("5. Playwright MCP як базовий шар", "course/05-playwright-mcp.md"),
                    ("6. Практика: повний цикл і порівняння «вручну vs агентами»", "course/06-practice-manual-vs-agents.md"),
                ],
            ),
            (
                "Частина II. Проєкт GreenKart (plan → generate → heal)",
                "Навчальний проєкт на застосунку GreenKart: плани тестування, ручні тест-кейси, "
                "автотести, знайдені дефекти та ролі агентів для DeepSeek Harness.",
                [
                    ("1. Огляд проєкту", "greenkart/README.md"),
                    ("2. План тестування (базовий)", "greenkart/specs/greenkart.plan.md"),
                    ("3. План групи Cart Edge Cases", "greenkart/specs/cart-edge-cases.plan.md"),
                    ("4. План групи Search and Cart", "greenkart/specs/search-and-cart.plan.md"),
                    ("5. План ручного тестування", "greenkart/test-cases/greenkart-test-plan.md"),
                    ("6. Ручні тест-кейси", None),
                    ("7. Звіт про дефект BUG-1: кошик втрачається при швидкому перезавантаженні",
                     "greenkart/BUG-1-cart-is-lost-on-immediate-reload.md"),
                    ("8. Звіт про дефект BUG-2: замовлення можна оформити з порожнім кошиком",
                     "greenkart/BUG-2-empty-cart-can-be-ordered.md"),
                    ("9. Ролі агентів (скіли для DSH)", None),
                    ("10. Визначення агентів Playwright (init-agents)", None),
                    ("11. Журнал проєкту (SESSION-NOTES)", "greenkart/SESSION-NOTES.md"),
                ],
            ),
            (
                "Частина III. Проєкт Conduit (Playwright MCP vs playwright-cli)",
                "Той самий тест-кейс автоматизовано двічі — через Playwright MCP і через playwright-cli — "
                "для порівняння підходів.",
                [
                    ("1. Огляд і порівняння підходів", "conduit/README.md"),
                    ("2. План тестування", "conduit/conduit.plan.md"),
                    ("3. Промпти генерації (MCP і CLI)", "conduit/prompts.txt"),
                    ("4. Журнал проєкту (SESSION-NOTES)", "conduit/SESSION-NOTES.md"),
                ],
            ),
        ],
        "groups": {
            "6. Ручні тест-кейси": ("greenkart/test-cases", "TC-"),
            "9. Ролі агентів (скіли для DSH)": ("greenkart/skills", None),
            "10. Визначення агентів Playwright (init-agents)": ("greenkart/agents", None),
        },
    },
}


def group_files(lang_root: str, spec) -> list[tuple[str, str]]:
    directory, prefix = spec
    root = os.path.join(lang_root, directory)
    if not os.path.isdir(root):
        return []
    names = sorted(
        name for name in os.listdir(root)
        if name.endswith(".md") and (prefix is None or name.startswith(prefix))
    )
    return [(name[:-3], os.path.join(root, name)) for name in names]


def build(lang: str) -> dict:
    labels = LABELS[lang]
    lang_root = os.path.join(I18N_ROOT, lang)
    output = os.path.join(WORKSPACE, f"Playwright-Test-Agents-Documentation-{lang.upper()}.docx")

    doc = builder.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(labels["title"])
    run.bold = True
    run.font.size = Pt(30)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(labels["subtitle"])
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run(labels["intro"])
    run.font.size = Pt(11)

    doc.add_paragraph()
    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = stamp.add_run(labels["compiled"].format(date=date.today().strftime("%d.%m.%Y")))
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_page_break()

    # ---- index
    doc.add_heading(labels["contents"], level=1)
    for part_title, _, documents in labels["parts"]:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(part_title)
        run.bold = True
        paragraph.paragraph_format.space_before = Pt(8)
        for label, relative in documents:
            if relative is None:
                entries = group_files(lang_root, labels["groups"][label])
                for sub_label, _sub_path in entries:
                    doc.add_paragraph(sub_label, style="List Bullet")
                if not entries:
                    doc.add_paragraph(label, style="List Bullet")
                continue
            doc.add_paragraph(label, style="List Bullet")

    stats = {"documents": 0}
    missing: list[str] = []

    for part_title, part_intro, documents in labels["parts"]:
        doc.add_page_break()
        doc.add_heading(part_title, level=1)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(part_intro)
        run.italic = True

        for label, relative in documents:
            if relative is None:
                group_number = label.split(".")[0]
                for sub_label, sub_path in group_files(lang_root, labels["groups"][label]):
                    doc.add_page_break()
                    doc.add_heading(f"{group_number}. {sub_label}", level=2)
                    with open(sub_path, encoding="utf-8") as handle:
                        builder.convert_markdown(doc, handle.read(), 2)
                    stats["documents"] += 1
                continue
            path = os.path.join(lang_root, relative)
            if not os.path.exists(path):
                missing.append(relative)
                continue
            doc.add_page_break()
            doc.add_heading(label, level=2)
            with open(path, encoding="utf-8") as handle:
                builder.convert_markdown(doc, handle.read(), 2)
            stats["documents"] += 1

    builder.add_page_number_footer(doc)
    doc.save(output)

    check = builder.Document(output)
    return {
        "output": output,
        "documents": stats["documents"],
        "paragraphs": len(check.paragraphs),
        "tables": len(check.tables),
        "words": sum(len(p.text.split()) for p in check.paragraphs),
        "size_kb": round(os.path.getsize(output) / 1024, 1),
        "missing": missing,
    }


if __name__ == "__main__":
    for language in (sys.argv[1:] or ["en", "uk"]):
        result = build(language)
        print(f"[{language.upper()}] {result['output']}")
        print(f"  документов: {result['documents']}, абзацев: {result['paragraphs']}, "
              f"таблиц: {result['tables']}, слов: {result['words']}, размер: {result['size_kb']} КБ")
        if result["missing"]:
            print(f"  НЕ хватает {len(result['missing'])} файлов:")
            for relative in result["missing"]:
                print(f"    {relative}")
